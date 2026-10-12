"""Moneda seleccionable en la exportación e importación de productos."""

from __future__ import annotations

import json
import unittest
from io import BytesIO

from openpyxl import load_workbook

from app import create_app, db
from app.inventario_comercial.productos_excel import (
    excel_productos_catalogo,
    excel_productos_plantilla,
    importar_productos_desde_excel,
)
from app.models import Empresa, InvProducto


class TestProductosExcelCurrency(unittest.TestCase):
    def setUp(self):
        self.app = create_app("testing")
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.create_all()
        self.empresa = Empresa(
            razon_social="Importadora SAS",
            slug="importadora-sas",
            nit="900001",
            moneda="COP",
            monedas_activas_json=json.dumps(["COP", "USD"]),
        )
        db.session.add(self.empresa)
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.ctx.pop()

    def test_plantilla_usa_moneda_elegida(self):
        contenido, nombre = excel_productos_plantilla(self.empresa.id, moneda="USD")
        wb = load_workbook(BytesIO(contenido), data_only=True)
        headers = [cell.value for cell in wb["Productos"][1]]

        self.assertIn("Precio compra (USD)", headers)
        self.assertIn("Precio venta (USD)", headers)
        self.assertNotIn("Precio venta (COP)", headers)
        self.assertEqual(nombre, "plantilla_productos_usd.xlsx")
        self.assertIn("USD", wb["Instrucciones"]["A7"].value)
        wb.close()

    def test_importar_una_moneda_conserva_las_otras(self):
        producto = InvProducto(
            empresa_id=self.empresa.id,
            sku="SKU-1",
            nombre="Producto",
            precio_compra=10000,
            precio_venta=15000,
            precios_json=json.dumps({"COP": 15000, "USD": 4}),
        )
        db.session.add(producto)
        db.session.commit()

        contenido, _ = excel_productos_plantilla(self.empresa.id, moneda="USD")
        wb = load_workbook(BytesIO(contenido))
        ws = wb["Productos"]
        headers = [cell.value for cell in ws[1]]
        ws.delete_rows(2, 1)
        fila = [None] * len(headers)
        fila[headers.index("Referencia")] = "SKU-1"
        fila[headers.index("Nombre")] = "Producto actualizado"
        fila[headers.index("Precio venta (USD)")] = 8.5
        ws.append(fila)
        salida = BytesIO()
        wb.save(salida)
        wb.close()

        resultado = importar_productos_desde_excel(
            self.empresa.id, salida.getvalue(), nombre_archivo="productos_usd.xlsx"
        )
        db.session.commit()
        db.session.refresh(producto)

        self.assertEqual(resultado.actualizados, 1)
        self.assertEqual(producto.nombre, "Producto actualizado")
        self.assertEqual(producto.precio_compra, 10000)
        self.assertEqual(producto.precios_venta("COP")["COP"], 15000)
        self.assertEqual(producto.precios_venta("COP")["USD"], 8.5)

    def test_catalogo_indica_moneda_elegida(self):
        producto = InvProducto(
            empresa_id=self.empresa.id,
            sku="SKU-2",
            nombre="Segundo",
            precio_compra=20000,
            precio_venta=25000,
            precios_json=json.dumps({"COP": 25000, "USD": 6.25}),
        )
        db.session.add(producto)
        db.session.commit()

        contenido, nombre = excel_productos_catalogo(
            self.empresa.id, [producto], moneda="USD"
        )
        wb = load_workbook(BytesIO(contenido), data_only=True)
        ws = wb["Productos"]
        headers = [cell.value for cell in ws[7]]
        datos = [cell.value for cell in ws[8]]

        self.assertIn("Precio venta (USD)", headers)
        self.assertEqual(datos[headers.index("Precio venta (USD)")], 6.25)
        self.assertIsNone(datos[headers.index("Precio compra (USD)")])
        self.assertIn("USD", nombre)
        wb.close()

    def test_importar_activa_la_moneda_del_archivo(self):
        self.empresa.monedas_activas_json = json.dumps(["COP"])
        db.session.commit()
        contenido, _ = excel_productos_plantilla(self.empresa.id, moneda="USD")
        wb = load_workbook(BytesIO(contenido))
        ws = wb["Productos"]
        headers = [cell.value for cell in ws[1]]
        ws.delete_rows(2, 1)
        fila = [None] * len(headers)
        fila[headers.index("Referencia")] = "USD-1"
        fila[headers.index("Nombre")] = "Producto en dólares"
        fila[headers.index("Precio venta (USD)")] = 12.5
        ws.append(fila)
        salida = BytesIO()
        wb.save(salida)
        wb.close()

        resultado = importar_productos_desde_excel(
            self.empresa.id, salida.getvalue(), nombre_archivo="productos_usd.xlsx"
        )
        db.session.commit()
        producto = InvProducto.query.filter_by(sku="USD-1").one()

        self.assertEqual(resultado.monedas_activadas, ["USD"])
        self.assertEqual(self.empresa.monedas_activas, ["COP", "USD"])
        self.assertEqual(producto.precios_venta("COP")["USD"], 12.5)


if __name__ == "__main__":
    unittest.main()
